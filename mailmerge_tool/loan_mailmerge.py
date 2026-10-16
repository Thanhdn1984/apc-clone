#!/usr/bin/env python3
# -*- coding: utf-8 -*-
import argparse, csv, json, math, os, re, sys, tempfile, zipfile
from collections import OrderedDict
from copy import deepcopy
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook, Workbook
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_BREAK
from docx.shared import Pt

BLANK = (None, '')

def is_blank(v):
    if v is None: return True
    if isinstance(v, float) and math.isnan(v): return True
    return str(v).strip() == ''

def safe_name(s):
    s = str(s or 'khach_hang').strip()
    s = re.sub(r'[\\/:*?"<>|]+', '-', s)
    s = re.sub(r'\s+', ' ', s)
    return s[:120] or 'khach_hang'

def read_excel(path, sheet=None):
    wb = load_workbook(path, data_only=True)
    ws = wb[sheet] if sheet else wb[wb.sheetnames[0]]
    rows = list(ws.iter_rows(values_only=True))
    if not rows: raise ValueError('File Excel trống')
    headers = [str(h).strip() if h is not None else '' for h in rows[0]]
    if not any(headers): raise ValueError('Không đọc được dòng tiêu đề Excel')
    seen = {}
    fixed=[]
    for i,h in enumerate(headers):
        h = h or f'cot_{i+1}'
        if h in seen:
            seen[h]+=1; h=f'{h}_{seen[h]}'
        else: seen[h]=1
        fixed.append(h)
    data=[]
    for ridx, row in enumerate(rows[1:], start=2):
        item={fixed[i]: row[i] if i < len(row) else None for i in range(len(fixed))}
        if all(is_blank(v) for v in item.values()):
            continue
        item['_excel_row']=ridx
        data.append(item)
    return fixed, data

def split_and_group(rows, group_col, common_cols, repeat_cols):
    missing = [c for c in [group_col,*common_cols,*repeat_cols] if c not in (rows[0].keys() if rows else [])]
    if missing: raise ValueError('Thiếu cột: ' + ', '.join(missing))
    groups=OrderedDict(); errors=[]
    for r in rows:
        reasons=[]
        if is_blank(r.get(group_col)): reasons.append(f'thiếu cột nhóm {group_col}')
        for c in common_cols:
            if is_blank(r.get(c)): reasons.append(f'thiếu thông tin chung {c}')
        if reasons:
            bad=dict(r); bad['_error']='; '.join(reasons); errors.append(bad); continue
        key=str(r[group_col]).strip()
        if key not in groups:
            groups[key]={'key':key,'common':{c:r.get(c) for c in common_cols},'loans':[],'rows':[]}
        groups[key]['loans'].append({c:r.get(c) for c in repeat_cols})
        groups[key]['rows'].append(r.get('_excel_row'))
    return list(groups.values()), errors

def fmt(v):
    if v is None: return ''
    if isinstance(v, datetime): return v.strftime('%d/%m/%Y')
    if isinstance(v, (int,float)) and not isinstance(v,bool):
        if float(v).is_integer(): return f'{int(v):,}'.replace(',', '.')
        return f'{v:,.2f}'.replace(',', 'X').replace('.', ',').replace('X','.')
    return str(v)

def add_letter(doc, group, common_cols, repeat_cols, title, idx, total):
    if idx:
        doc.add_page_break()
    p=doc.add_paragraph()
    p.alignment=1
    run=p.add_run('AGRIBANK')
    run.bold=True; run.font.size=Pt(14)
    p=doc.add_paragraph(); p.alignment=1
    r=p.add_run(title); r.bold=True; r.font.size=Pt(13)
    doc.add_paragraph('')
    for c in common_cols:
        doc.add_paragraph(f'{c}: {fmt(group["common"].get(c))}')
    doc.add_paragraph('')
    table=doc.add_table(rows=1, cols=len(repeat_cols)+1)
    table.style='Table Grid'
    hdr=table.rows[0].cells
    hdr[0].text='STT'
    for i,c in enumerate(repeat_cols, start=1): hdr[i].text=c
    for n, loan in enumerate(group['loans'], start=1):
        cells=table.add_row().cells; cells[0].text=str(n)
        for i,c in enumerate(repeat_cols, start=1): cells[i].text=fmt(loan.get(c))
    doc.add_paragraph('')
    doc.add_paragraph(f'Tổng số khoản vay: {len(group["loans"])}')

def write_errors_xlsx(errors, out):
    wb=Workbook(); ws=wb.active; ws.title='Dong loi'
    if not errors:
        ws.append(['Không có dòng lỗi']); wb.save(out); return
    headers=['_excel_row','_error']+[h for h in errors[0].keys() if h not in ('_excel_row','_error')]
    ws.append(headers)
    for e in errors: ws.append([e.get(h) for h in headers])
    wb.save(out)

def generate(excel, out_docx, group_col, common_cols, repeat_cols, sheet=None, title='THƯ XÁC NHẬN SỐ DƯ TÀI KHOẢN', separate_dir=None, errors_xlsx=None):
    headers, rows=read_excel(excel, sheet)
    if not rows: raise ValueError('Không có dòng dữ liệu sau tiêu đề')
    groups, errors=split_and_group(rows, group_col, common_cols, repeat_cols)
    if not groups: raise ValueError('Không có khách hàng hợp lệ để xuất thư')
    if separate_dir:
        Path(separate_dir).mkdir(parents=True, exist_ok=True)
        for i,g in enumerate(groups):
            d=Document(); add_letter(d,g,common_cols,repeat_cols,title,0,len(groups))
            d.save(Path(separate_dir)/f'{safe_name(g["key"])}.docx')
    else:
        doc=Document()
        for i,g in enumerate(groups): add_letter(doc,g,common_cols,repeat_cols,title,i,len(groups))
        doc.save(out_docx)
    if errors_xlsx: write_errors_xlsx(errors, errors_xlsx)
    return {'input_rows':len(rows),'letters':len(groups),'error_rows':len(errors),'groups':[{ 'key':g['key'], 'loan_count':len(g['loans']), 'excel_rows':g['rows']} for g in groups]}

def parse_cols(s): return [x.strip() for x in (s or '').split(',') if x.strip()]

def main(argv=None):
    ap=argparse.ArgumentParser(description='Mail Merge khoản vay: gộp nhiều dòng Excel thành 1 thư/khách hàng')
    ap.add_argument('excel', help='file .xlsx')
    ap.add_argument('--list-columns', action='store_true')
    ap.add_argument('--sheet')
    ap.add_argument('--group-col')
    ap.add_argument('--common-cols', help='cách nhau bằng dấu phẩy')
    ap.add_argument('--repeat-cols', help='cách nhau bằng dấu phẩy')
    ap.add_argument('--out', default='thu_xac_nhan.docx')
    ap.add_argument('--errors-out', default='dong_loi.xlsx')
    ap.add_argument('--separate-dir')
    ap.add_argument('--preview', action='store_true')
    args=ap.parse_args(argv)
    headers, rows=read_excel(args.excel, args.sheet)
    if args.list_columns:
        print(json.dumps({'columns':headers,'rows':len(rows)}, ensure_ascii=False, indent=2)); return 0
    if not args.group_col or not args.common_cols or not args.repeat_cols:
        ap.error('cần --group-col --common-cols --repeat-cols hoặc dùng --list-columns')
    groups, errors=split_and_group(rows,args.group_col,parse_cols(args.common_cols),parse_cols(args.repeat_cols))
    summary={'input_rows':len(rows),'letters':len(groups),'error_rows':len(errors),'groups':[{'key':g['key'],'loan_count':len(g['loans']),'excel_rows':g['rows']} for g in groups]}
    print(json.dumps(summary, ensure_ascii=False, indent=2))
    if args.preview: return 0
    res=generate(args.excel,args.out,args.group_col,parse_cols(args.common_cols),parse_cols(args.repeat_cols),args.sheet,separate_dir=args.separate_dir,errors_xlsx=args.errors_out)
    print('DONE:', json.dumps(res, ensure_ascii=False))
    return 0
if __name__=='__main__': raise SystemExit(main())
