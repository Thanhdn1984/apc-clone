import json, subprocess, sys, zipfile
from pathlib import Path
from openpyxl import Workbook, load_workbook

ROOT=Path(__file__).resolve().parents[1]
CLI=ROOT/'mailmerge_tool'/'loan_mailmerge.py'

def make_sample(path):
    wb=Workbook(); ws=wb.active
    ws.append(['brcd','custnm','addr1','apprseq','Sum','ccy'])
    ws.append(['KH001','Nguyễn Văn A','Tây Giang','HD001',1000000,'VND'])
    ws.append(['KH001','Nguyễn Văn A','Tây Giang','HD002',2000000,'VND'])
    ws.append(['KH001','Nguyễn Văn A','Tây Giang','HD003',3000000,'VND'])
    ws.append(['KH002','Trần Thị B','Đà Nẵng','HD004',4000000,'VND'])
    ws.append(['KH003','', 'Quảng Nam','HD005',5000000,'VND'])
    wb.save(path)

def test_group_and_error(tmp_path):
    x=tmp_path/'sample.xlsx'; out=tmp_path/'out.docx'; err=tmp_path/'err.xlsx'
    make_sample(x)
    p=subprocess.run([sys.executable,str(CLI),str(x),'--group-col','brcd','--common-cols','brcd,custnm,addr1','--repeat-cols','apprseq,Sum,ccy','--out',str(out),'--errors-out',str(err)],capture_output=True,text=True,check=True)
    assert '"letters": 2' in p.stdout
    assert '"error_rows": 1' in p.stdout
    assert out.exists() and err.exists()
    wb=load_workbook(err); ws=wb.active
    assert ws.max_row==2
    xml=zipfile.ZipFile(out).read('word/document.xml').decode('utf-8')
    assert 'Nguyễn Văn A' in xml and 'HD001' in xml and 'HD002' in xml and 'HD003' in xml
    assert 'Trần Thị B' in xml and 'HD004' in xml
    assert 'HD005' not in xml
